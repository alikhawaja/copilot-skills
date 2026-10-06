"""Check a staged monthly services workbook against independently retrieved cost totals."""

import argparse
from collections import defaultdict
from contextlib import closing
from decimal import Decimal, InvalidOperation
from pathlib import Path
from zipfile import BadZipFile

from openpyxl import load_workbook


SHEETS = [
    "Services",
    "Basis and reconciliation",
    "Attribution evidence",
    "SKU evidence",
    "Subscription reconciliation",
    "Unallocated costs",
]
TOLERANCE = Decimal("0.01")


def require(condition, message):
    if not condition:
        raise ValueError(message)


def amount(value):
    require(isinstance(value, (int, float, Decimal)) and not isinstance(value, bool),
            f"Cost must be numeric, got {value!r}")
    return Decimal(str(value))


def rows(sheet, headers):
    data = sheet.values
    require(next(data, None) == tuple(headers), f"{sheet.title}: unexpected headers")
    return list(data)


def close(actual, expected, description):
    require(abs(actual - expected) <= TOLERANCE,
            f"{description}: {actual} does not reconcile to {expected}")


def validate(path, month, subscriptions, expected_total):
    require(path.is_file(), f"Missing services workbook: {path}")
    with closing(load_workbook(path, read_only=True, data_only=True)) as book:
        require(book.sheetnames == SHEETS, f"Expected worksheets in order: {SHEETS}")
        services = rows(book["Services"], ["ServiceName", "Workload", "ACR"])
        basis = dict((r[0], r[1]) for r in rows(
            book["Basis and reconciliation"], ["Metric", "Value", "Meaning"]))
        require(basis.get("Reporting month") == month, "Reporting month mismatch")
        require(basis.get("Measure") == "ActualCost", "Expected historical ActualCost")
        currency = basis.get("Reporting currency")
        require(isinstance(currency, str) and currency.strip(), "Missing reporting currency")
        require(basis.get("Source currencies"), "Missing source currencies")
        require(basis.get("UTC retrieval time"), "Missing UTC retrieval time")
        require(basis.get("Exchange rate") is not None, "Missing conversion basis")
        cost = f"ActualCost {currency} equivalent"
        evidence = rows(book["Attribution evidence"], [
            "Workload", "ServiceName", "Source subscription", "Subscription ID",
            "Source currency", cost, "Attribution basis",
        ])
        sku = rows(book["SKU evidence"], [
            "Workload", "ServiceName", "Subscription ID", "Source resource ID",
            "Source resource type", "Discovered source SKU", "SKU provenance",
            "Source currency", cost,
        ])
        reconciliation = rows(book["Subscription reconciliation"], [
            "Subscription", "Subscription ID", "Source currency", cost,
            f"Attributed {currency}", f"Unallocated {currency}", "Note",
        ])
        unallocated = rows(book["Unallocated costs"], [
            "Subscription", "Subscription ID", "ServiceName", "Source currency",
            cost, "Reason",
        ])

        service_totals = defaultdict(Decimal)
        for service, workload, value in services:
            require(service and workload, "Empty Services label or workload")
            service_totals[(workload, service)] += amount(value)
        attributed = defaultdict(Decimal)
        per_subscription = defaultdict(Decimal)
        for workload, service, _, sid, source_currency, value, reason in evidence:
            require(workload and service and source_currency and reason,
                    "Incomplete attribution evidence")
            require(sid in subscriptions, f"Unexpected attribution subscription: {sid}")
            attributed[(workload, service)] += amount(value)
            per_subscription[sid] += amount(value)
        require(set(service_totals) == set(attributed),
                "Services and Attribution evidence labels differ")
        for key, value in service_totals.items():
            close(value, attributed[key], f"Services {key}")

        labelled = {key: value for key, value in service_totals.items() if ": " in key[1]}
        sku_totals = defaultdict(Decimal)
        for workload, service, sid, resource_id, resource_type, tier, provenance, source_currency, value in sku:
            require(sid in subscriptions and resource_id and resource_type and tier
                    and provenance and source_currency, "Incomplete SKU evidence")
            require(": " in service and service.endswith(f": {tier}"),
                    f"SKU evidence does not match service label: {service}")
            sku_totals[(workload, service)] += amount(value)
        require(set(labelled) == set(sku_totals), "SKU-labelled services and SKU evidence differ")
        for key, value in labelled.items():
            close(value, sku_totals[key], f"SKU evidence {key}")

        missing = defaultdict(Decimal)
        for _, sid, service, source_currency, value, reason in unallocated:
            require(sid in subscriptions and service and source_currency and reason,
                    "Incomplete unallocated-cost reason")
            missing[sid] += amount(value)
        retrieved = Decimal(0)
        seen = set()
        for name, sid, source_currency, total, allocated, remaining, note in reconciliation:
            require(name and source_currency and note, "Incomplete subscription reconciliation")
            require(sid in subscriptions and sid not in seen,
                    f"Unexpected or duplicate subscription reconciliation: {sid}")
            seen.add(sid)
            close(amount(allocated), per_subscription[sid], f"Attributed {sid}")
            close(amount(remaining), missing[sid], f"Unallocated {sid}")
            close(amount(total), amount(allocated) + amount(remaining), f"Total {sid}")
            retrieved += amount(total)
        require(seen == subscriptions, f"Missing subscription reconciliation: {subscriptions - seen}")
        close(retrieved, expected_total, "Independent ActualCost total")
        close(sum(service_totals.values()) + sum(missing.values()), expected_total,
              "Services plus unallocated costs")
    print(f"Validated {path}: {month}, {currency} {retrieved}, {len(subscriptions)} subscriptions")


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--month", required=True, help="Confirmed YYYY-MM billing month")
    parser.add_argument("--subscription", action="append", required=True, help="Selected subscription ID")
    parser.add_argument("--total", required=True, help="Independent ActualCost total in reporting currency")
    args = parser.parse_args()
    try:
        validate(args.workbook, args.month, set(args.subscription), Decimal(args.total))
    except BadZipFile:
        parser.exit(1, "Services workbook validation failed: file is not a standard XLSX; "
                    "if Office protection transformed the published copy, validate the staged "
                    "workbook and inspect the published copy in Excel.\n")
    except (ValueError, InvalidOperation, OSError) as exc:
        parser.exit(1, f"Services workbook validation failed: {exc}\n")


if __name__ == "__main__":
    main()
